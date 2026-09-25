"""Document loading and text helpers shared by the parsers.

PDFs are read with pdfplumber at word level so every word keeps its x
position: BoursoBank statements print debits and credits in two columns
with nothing but position to tell them apart, which plain text
extraction loses.
"""

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

# BoursoBank's statement fonts map accented glyphs to the wrong code points
# (é -> Ø, è -> Ł) and leave some glyphs as raw CIDs.
GLYPH_FIXES = {
    "Ø": "é", "Ł": "è",
    "(cid:128)": "€", "(cid:224)": "à", "(cid:176)": "°", "(cid:150)": "-",
    "’": "'", " ": " ", " ": " ",
}

DATE = r"(\d{2}/\d{2}/\d{4})"
# French amounts: "1.873,44", "3 614,91", "500,00", "-12,30"
FR_AMOUNT = r"-?\d{1,3}(?:[. ]\d{3})*,\d{2}"
FR_NUMBER = r"-?\d{1,3}(?:[. ]\d{3})*,\d+|-?\d+,\d+"

FR_MONTHS = {
    "janvier": 1, "février": 2, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5,
    "juin": 6, "juillet": 7, "août": 8, "aout": 8, "septembre": 9,
    "octobre": 10, "novembre": 11, "décembre": 12, "decembre": 12,
}


def fix_glyphs(text):
    for old, new in GLYPH_FIXES.items():
        text = text.replace(old, new)
    return text


def fr_decimal(token):
    """"1.873,44" / "3 614,91" / "0,74627" -> Decimal."""
    return Decimal(token.replace(" ", "").replace(".", "").replace(",", "."))


def fr_date(token):
    day, month, year = (int(p) for p in token.split("/"))
    return date(year, month, day)


def fr_long_date(text):
    """"20 mars 2026" -> date(2026, 3, 20)."""
    m = re.search(r"(\d{1,2})\s+([a-zéû]+)\s+(\d{4})", text, re.I)
    if not m or m.group(2).lower() not in FR_MONTHS:
        return None
    return date(int(m.group(3)), FR_MONTHS[m.group(2).lower()], int(m.group(1)))


@dataclass
class Word:
    x0: float
    x1: float
    text: str


@dataclass
class Line:
    top: float
    words: list

    @property
    def text(self):
        return " ".join(w.text for w in self.words)


@dataclass
class RawDocument:
    """What a parser receives: the path, the full text, and — for PDFs —
    pages of positioned lines; for CSVs, the rows; for XLSX workbooks,
    the visible sheets."""

    path: str
    text: str
    pages: list = field(default_factory=list)  # list[list[Line]]
    rows: Optional[list] = None  # CSV rows as dicts
    sheets: Optional[dict] = None  # XLSX: sheet name -> list of rows (lists of cell values)

    @property
    def name(self):
        return Path(self.path).name


def group_lines(words, y_tolerance=3.0):
    """Cluster words into lines by vertical position. A tolerance rather
    than exact equality: a trailing amount is sometimes printed a point
    or two off its label's baseline."""
    lines = []
    # Rotated text (a letter printed sideways in the margin) is not part of
    # any table row, but its top can land on a row's baseline and glue a
    # stray "V" onto a fund name.
    words = [w for w in words if w.get("upright", True)]
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        word = Word(w["x0"], w["x1"], fix_glyphs(w["text"]))
        if lines and abs(w["top"] - lines[-1].top) <= y_tolerance:
            lines[-1].words.append(word)
        else:
            lines.append(Line(w["top"], [word]))
    for line in lines:
        line.words.sort(key=lambda w: w.x0)
    return lines


def cell_value(value):
    """An openpyxl cell value as the parsers want it: floats and ints ->
    Decimal (through str, so -33.91 stays -33.91), datetime -> date,
    blank strings -> None, anything else unchanged."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str) and not value.strip():
        return None
    return value


def load_workbook(path):
    """Visible sheets only: exports keep stale template sheets around as
    'veryHidden', and those must never be read as accounts."""
    import openpyxl  # imported lazily: callers without XLSX files don't need it

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {
            ws.title: [[cell_value(v) for v in row] for row in ws.iter_rows(values_only=True)]
            for ws in workbook.worksheets
            if ws.sheet_state == "visible"
        }
    finally:
        workbook.close()
    return sheets


def sheets_text(sheets):
    """Sheet names and every text cell, one per line, for detect()."""
    lines = []
    for name, rows in sheets.items():
        lines.append(name)
        for row in rows:
            seen = None
            for value in row:
                # exports repeat a merged title in every cell of the row
                if isinstance(value, str) and value != seen:
                    lines.append(value)
                    seen = value
    return "\n".join(lines)


def load_document(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        sheets = load_workbook(path)
        return RawDocument(str(path), sheets_text(sheets), sheets=sheets)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        header = ",".join(rows[0].keys()) if rows else ""
        return RawDocument(str(path), header, rows=rows)

    import pdfplumber  # imported lazily: CSV-only callers don't need it

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(group_lines(page.extract_words(x_tolerance=1)))
    text = "\n".join(line.text for page in pages for line in page)
    return RawDocument(str(path), text, pages=pages)


def column_edges(line, labels):
    """Right edge of each column header on a header line, e.g.
    {"debit": 465.0, "credit": 533.0}. A header like "Débit EUR" is two
    words; the column's right edge is the last word's."""
    edges = {}
    words = line.words
    for i, w in enumerate(words):
        for key, pattern in labels.items():
            if re.fullmatch(pattern, w.text):
                right = w.x1
                if i + 1 < len(words) and words[i + 1].text == "EUR" and words[i + 1].x0 - w.x1 < 6:
                    right = words[i + 1].x1
                edges[key] = right
    return edges


def classify_amount(word, edges, tolerance=25.0):
    """Which column a right-aligned amount belongs to: the header whose
    right edge is nearest, if within tolerance."""
    best = min(edges.items(), key=lambda kv: abs(kv[1] - word.x1), default=None)
    if best is None or abs(best[1] - word.x1) > tolerance:
        return None
    return best[0]
